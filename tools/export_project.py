import argparse
from datetime import date, datetime, time
from hashlib import sha256
from html import escape
import json
from pathlib import Path
import warnings

import openpyxl
from openpyxl.utils import get_column_letter


def text(value):
    if value is None:
        return ""
    if isinstance(value, (date, datetime, time)):
        value = value.isoformat()
    return escape(str(value), quote=False).replace("|", "&#124;").replace("\r\n", "\n").replace("\n", "<br>")


def export_sheet(sheet, cached):
    rows = [list(row) for row in sheet.iter_rows()]
    populated = [cell for row in rows for cell in row if cell.value is not None]
    header_row = {"Renovatie": 7, "Planning": 5}.get(sheet.title, 1)
    columns = sorted({cell.column for cell in populated if cell.row >= header_row})
    lines = [f"# {sheet.title}", "", "Bedragen in EUR tenzij anders vermeld. Rij- en kolomcodes verwijzen naar het bewaarde Excelbestand.", "", "Dit is een statische momentopname: formules rekenen niet in Markdown. Leeg is onbekend/niet ingevuld, niet nul. Formulewaarden zijn caches; geen volledige Excel-herberekening uitgevoerd.", ""]
    summary = [cell for cell in populated if cell.row < header_row]
    if summary:
        lines += ["## Samenvatting", "", "| Cel | Waarde |", "| --- | --- |"]
        lines += [f"| {cell.coordinate} | {text(cached[cell.coordinate].value if cell.data_type == 'f' else cell.value)} |" for cell in summary]
        lines += [""]
    headers = [f"{get_column_letter(column)}: {text(sheet.cell(header_row, column).value) or 'Zonder titel'}" for column in columns]
    lines += ["## Gegevens", "", "| Rij | " + " | ".join(headers) + " |", "| --- | " + " | ".join("---" for column in columns) + " |"]
    exported_rows = 0
    for row in rows[header_row:]:
        if not any(cell.value is not None for cell in row):
            continue
        values = []
        for column in columns:
            cell = sheet.cell(row[0].row, column)
            value = cached[cell.coordinate].value if cell.data_type == "f" else cell.value
            values.append(text(value))
        lines.append(f"| {row[0].row} | " + " | ".join(values) + " |")
        exported_rows += 1
    formulas = [cell for cell in populated if cell.data_type == "f"]
    if formulas:
        lines += ["", "## Formules (bronbestand)", "", "| Cel | Formule | Gecachete waarde |", "| --- | --- | --- |"]
        lines += [f"| {cell.coordinate} | {text(cell.value)} | {text(cached[cell.coordinate].value)} |" for cell in formulas]
    annotations = [cell for cell in populated if cell.comment or cell.hyperlink]
    if annotations:
        lines += ["", "## Opmerkingen en koppelingen", "", "| Cel | Opmerking | Koppeling |", "| --- | --- | --- |"]
        lines += [f"| {cell.coordinate} | {text(cell.comment.text if cell.comment else None)} | {text((cell.hyperlink.target or cell.hyperlink.location) if cell.hyperlink else None)} |" for cell in annotations]
    if sheet.merged_cells.ranges:
        lines += ["", "Samengevoegde broncellen: " + ", ".join(str(area) for area in sheet.merged_cells.ranges) + "."]
    if sheet._charts:
        lines += ["", "De Excelgrafiek is niet overgenomen; de tabel en onderliggende formules staan hierboven."]
    document = "\n".join(lines) + "\n"
    for cell in populated:
        value = cached[cell.coordinate].value if cell.data_type == "f" else cell.value
        assert not value or text(value) in document, f"Missing {sheet.title}!{cell.coordinate}"
    for cell in formulas:
        assert f"| {cell.coordinate} | {text(cell.value)} |" in document
    assert len([line for line in document.splitlines() if line.startswith("| ") and line.split(" | ")[0][2:].isdigit()]) == exported_rows
    return document, {"sheet": sheet.title, "data_rows": exported_rows, "populated_cells": len(populated), "formulas": len(formulas)}


def main():
    parser = argparse.ArgumentParser(description="Export a workbook to one Markdown file per sheet without modifying the source.")
    parser.add_argument("workbook", type=Path)
    parser.add_argument("destination", type=Path)
    arguments = parser.parse_args()
    original_hash = sha256(arguments.workbook.read_bytes()).hexdigest()
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Data Validation extension is not supported.*")
        workbook = openpyxl.load_workbook(arguments.workbook, data_only=False)
        cached = openpyxl.load_workbook(arguments.workbook, data_only=True)
    documents = []
    for sheet in workbook:
        assert Path(sheet.title).name == sheet.title and not any(character in sheet.title for character in '<>:"/\\|?*'), "Unsafe sheet filename"
        path = arguments.destination / f"{sheet.title}.md"
        if path.exists():
            raise FileExistsError(f"Refusing to overwrite editable Markdown: {path}")
        document, counts = export_sheet(sheet, cached[sheet.title])
        documents.append((path, document, counts))
    arguments.destination.mkdir(parents=True, exist_ok=True)
    for path, document, counts in documents:
        path.write_text(document, encoding="utf-8")
        assert path.read_text(encoding="utf-8") == document
    workbook.close()
    cached.close()
    assert sha256(arguments.workbook.read_bytes()).hexdigest() == original_hash, "Source workbook changed"
    print(json.dumps({"source_sha256": original_hash, "sheets": [counts for path, document, counts in documents]}, indent=2))


if __name__ == "__main__":
    assert text("a|b\nc<d") == "a&#124;b<br>c&lt;d"
    assert text(0) == "0"
    assert text(None) == ""
    assert text(date(2026, 10, 11)) == "2026-10-11"
    main()